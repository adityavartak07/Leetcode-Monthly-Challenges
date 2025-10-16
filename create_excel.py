#!/usr/bin/env python3
"""
Script to create an Excel file and write numbers in it.
"""

from openpyxl import Workbook

def create_excel_with_numbers():
    """Create an Excel file with numbers."""
    # Create a new workbook and select the active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Numbers"
    
    # Write numbers in the first column (1 to 100)
    ws['A1'] = 'Number'  # Header
    for i in range(1, 101):
        ws[f'A{i+1}'] = i
    
    # Write numbers in a grid pattern (10x10)
    ws['C1'] = 'Grid of Numbers'
    for row in range(1, 11):
        for col in range(1, 11):
            ws.cell(row=row+1, column=col+2, value=(row-1)*10 + col)
    
    # Save the workbook
    filename = 'numbers.xlsx'
    wb.save(filename)
    print(f"Excel file '{filename}' created successfully!")
    print(f"- Column A contains numbers 1 to 100")
    print(f"- Columns C-L contain a 10x10 grid of numbers 1 to 100")

if __name__ == "__main__":
    create_excel_with_numbers()
